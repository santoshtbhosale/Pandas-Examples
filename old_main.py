import customtkinter as ctk
from tkinter import messagebox, filedialog
from tkcalendar import DateEntry
from datetime import datetime
import sqlite3
import os

# PDF Generation Imports
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors

# Excel Export Imports
import openpyxl

ctk.set_appearance_mode("light")
ctk.set_default_color_theme("blue")

# Global Variables / Data Stores
project_data = {}
project_id = "WD-" + datetime.now().strftime("%Y%m%d-%H%M%S")
residential_data = []
commercial_data = []
other_data = {}
calculated_results = {"Plot-A": {}, "Plot-B": {}, "Total": {}}

# ----------------- DATABASE SETUP -----------------
def init_db():
    conn = sqlite3.connect("water_demand.db")
    cursor = conn.cursor()
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS projects (
            project_id TEXT PRIMARY KEY,
            project_name TEXT,
            client_name TEXT,
            total_water_demand REAL,
            stp_capacity_a REAL,
            stp_capacity_b REAL,
            date TEXT
        )
    """)
    conn.commit()
    conn.close()

init_db()

app = ctk.CTk()
app.title("Water Demand Report Generator - Advanced 8-Page Edition")
app.geometry("1050x780")
app.resizable(False, False)

container = ctk.CTkFrame(app, fg_color="transparent")
container.pack(fill="both", expand=True, padx=10, pady=10)

frames = {}

def show_frame(page_name):
    frame = frames[page_name]
    frame.tkraise()

# ----------------- CALCULATIONS ENGINE -----------------
def perform_calculations():
    global calculated_results
    
    plots = ["Plot-A", "Plot-B"]
    for plot in plots:
        # Residential calculations
        res_plot = [r for r in residential_data if r["Plot"] == plot]
        res_pop = sum(r["Population"] for r in res_plot)
        res_dom = res_pop * 105
        res_flu = res_pop * 30
        
        # Commercial calculations
        com_plot = [c for c in commercial_data if c["Plot"] == plot]
        com_pop = sum(c["Auto Population"] for c in com_plot)
        com_dom = com_pop * 25
        com_flu = com_pop * 20
        
        total_water = (res_dom + res_flu) + (com_dom + com_flu)
        sewage_gen = ((res_dom + com_dom) + (res_flu + com_flu)) * 0.90
        stp_cap = round(sewage_gen / 1000, 2)
        
        # Tank Sizing
        ugt_dom = (res_dom + com_dom) * 2  # 2 days storage
        ugt_flu = (res_flu + com_flu) * 1  # 1 day storage
        fire_tank = other_data.get(f"Fire Tank {plot}", 200000)
        
        calculated_results[plot] = {
            "Population": res_pop + com_pop,
            "Res Pop": res_pop,
            "Com Pop": com_pop,
            "Total Water (LPD)": total_water,
            "Sewage Gen (LPD)": sewage_gen,
            "STP Capacity (KLD)": stp_cap,
            "UGT Domestic (Liters)": ugt_dom,
            "UGT Flushing (Liters)": ugt_flu,
            "Fire Tank (Liters)": fire_tank
        }

    # Grand Totals
    calculated_results["Total"] = {
        "Total Population": calculated_results["Plot-A"]["Population"] + calculated_results["Plot-B"]["Population"],
        "Total Water (LPD)": calculated_results["Plot-A"]["Total Water (LPD)"] + calculated_results["Plot-B"]["Total Water (LPD)"],
        "Total STP Capacity (KLD)": calculated_results["Plot-A"]["STP Capacity (KLD)"] + calculated_results["Plot-B"]["STP Capacity (KLD)"]
    }

# ----------------- PDF GENERATOR (8-PAGE FORMAT) -----------------
def export_pdf():
    file_path = filedialog.asksaveasfilename(defaultextension=".pdf", filetypes=[("PDF files", "*.pdf")], initialfile="8_Page_Water_Demand_Report.pdf")
    if not file_path:
        return
    try:
        doc = SimpleDocTemplate(file_path, pagesize=letter, rightMargin=20, leftMargin=20, topMargin=20, bottomMargin=20)
        story = []
        styles = getSampleStyleSheet()
        
        th_style = ParagraphStyle('TH', parent=styles['Normal'], fontSize=6.5, textColor=colors.whitesmoke, alignment=1, fontName='Helvetica-Bold')
        tc_style = ParagraphStyle('TC', parent=styles['Normal'], fontSize=6, textColor=colors.HexColor("#111111"), alignment=1)
        tc_left = ParagraphStyle('TCL', parent=styles['Normal'], fontSize=6, textColor=colors.HexColor("#111111"), alignment=0)
        sec_style = ParagraphStyle('Sec', parent=styles['Normal'], fontSize=8, textColor=colors.whitesmoke, alignment=1, fontName='Helvetica-Bold')
        
        footer_p = Paragraph("American Edge Engineers Pvt. Ltd. | Austin | New York | Pune | Info@americanedgeee.com", ParagraphStyle('Foot', fontSize=6, alignment=1))

        eng_header = Paragraph(f"<b>DESIGN ENGINEER: {project_data.get('Engineer Name', '').upper()} &nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp; DATE: {project_data.get('Date', '')}</b>", ParagraphStyle('Eng', fontSize=7, fontName='Helvetica-Bold'))

        # --- PAGE 1: COVER ---
        story.append(Spacer(1, 150))
        story.append(Paragraph("<b>AMERICAN EDGE ENGINEERS PVT. LTD.</b>", ParagraphStyle('H1', fontSize=16, alignment=1)))
        story.append(Spacer(1, 40))
        cover_meta = [
            [Paragraph("<b>TITLE</b>", tc_left), Paragraph(": WATER DEMAND CALCULATION", tc_left)],
            [Paragraph("<b>PROJECT NAME</b>", tc_left), Paragraph(f": {project_data.get('Project Name', '')}", tc_left)],
            [Paragraph("<b>CLIENT NAME</b>", tc_left), Paragraph(f": {project_data.get('Client Name', '')}", tc_left)],
        ]
        t_cover = Table(cover_meta, colWidths=[120, 400])
        t_cover.setStyle(TableStyle([('GRID', (0,0), (-1,-1), 1, colors.black), ('PADDING', (0,0), (-1,-1), 8)]))
        story.append(t_cover)
        story.append(PageBreak())

        # --- PAGE 2: CONSOLIDATED STATEMENT ---
        story.append(eng_header)
        story.append(Spacer(1, 5))
        story.append(Table([[Paragraph("<b>CONSOLIDATED STATEMENT</b>", sec_style)]], colWidths=[554], style=[('BACKGROUND', (0,0), (-1,-1), colors.HexColor("#34495E"))]))
        story.append(Spacer(1, 5))
        con_data = [
            [Paragraph("<b>DESCRIPTION</b>", th_style), Paragraph("<b>PLOT-A</b>", th_style), Paragraph("<b>PLOT-B</b>", th_style), Paragraph("<b>TOTAL</b>", th_style)],
            [Paragraph("Total Residential Population", tc_left), Paragraph(str(calculated_results["Plot-A"]["Res Pop"]), tc_style), Paragraph(str(calculated_results["Plot-B"]["Res Pop"]), tc_style), Paragraph(str(calculated_results["Plot-A"]["Res Pop"]+calculated_results["Plot-B"]["Res Pop"]), tc_style)],
            [Paragraph("Total Water Demand (LPD)", tc_left), Paragraph(str(calculated_results["Plot-A"]["Total Water (LPD)"]), tc_style), Paragraph(str(calculated_results["Plot-B"]["Total Water (LPD)"]), tc_style), Paragraph(str(calculated_results["Total"]["Total Water (LPD)"]), tc_style)],
            [Paragraph("Total STP Capacity (KLD)", tc_left), Paragraph(str(calculated_results["Plot-A"]["STP Capacity (KLD)"]), tc_style), Paragraph(str(calculated_results["Plot-B"]["STP Capacity (KLD)"]), tc_style), Paragraph(str(calculated_results["Total"]["Total STP Capacity (KLD)"]), tc_style)],
        ]
        t_con = Table(con_data, colWidths=[200, 100, 100, 100])
        t_con.setStyle(TableStyle([('GRID', (0,0), (-1,-1), 0.5, colors.black), ('BACKGROUND', (0,0), (-1,0), colors.HexColor("#7F8C8D"))]))
        story.append(t_con)
        story.append(PageBreak())

        # Function to generate Plot Demand Page
        def make_plot_demand_page(plot_name):
            story.append(eng_header)
            story.append(Spacer(1, 5))
            story.append(Table([[Paragraph(f"<b>WATER DEMAND - {plot_name}</b>", sec_style)]], colWidths=[554], style=[('BACKGROUND', (0,0), (-1,-1), colors.HexColor("#E67E22"))]))
            story.append(Spacer(1, 5))
            
            res_p = [r for r in residential_data if r["Plot"] == plot_name]
            r_data = [[Paragraph("<b>SR.NO</b>", th_style), Paragraph("<b>WING</b>", th_style), Paragraph("<b>FLATS</b>", th_style), Paragraph("<b>POPULATION</b>", th_style), Paragraph("<b>DOMESTIC</b>", th_style), Paragraph("<b>FLUSHING</b>", th_style), Paragraph("<b>TOTAL</b>", th_style)]]
            for idx, r in enumerate(res_p, 1):
                r_data.append([Paragraph(str(idx), tc_style), Paragraph(r["Wing"], tc_style), Paragraph(str(r["No. Of Flats"]), tc_style), Paragraph(str(r["Population"]), tc_style), Paragraph(str(r["Population"]*105), tc_style), Paragraph(str(r["Population"]*30), tc_style), Paragraph(str(r["Population"]*135), tc_style)])
            
            t_res = Table(r_data, colWidths=[40, 80, 60, 80, 90, 90, 90])
            t_res.setStyle(TableStyle([('GRID', (0,0), (-1,-1), 0.5, colors.black), ('BACKGROUND', (0,0), (-1,0), colors.HexColor("#34495E"))]))
            story.append(t_res)
            story.append(PageBreak())

        # --- PAGE 3 & 4: PLOT A & B DEMAND ---
        make_plot_demand_page("Plot-A")
        make_plot_demand_page("Plot-B")

        # Function to generate UGT Page
        def make_ugt_page(plot_name):
            story.append(eng_header)
            story.append(Spacer(1, 5))
            story.append(Table([[Paragraph(f"<b>UGT & OHT DETAILS - {plot_name}</b>", sec_style)]], colWidths=[554], style=[('BACKGROUND', (0,0), (-1,-1), colors.HexColor("#16A085"))]))
            story.append(Spacer(1, 5))
            u_data = [
                [Paragraph("<b>DESCRIPTION</b>", th_style), Paragraph("<b>STORAGE (DAYS)</b>", th_style), Paragraph("<b>TOTAL STORAGE (LITERS)</b>", th_style)],
                [Paragraph("DOMESTIC WATER TANK", tc_left), Paragraph("2", tc_style), Paragraph(str(calculated_results[plot_name]["UGT Domestic (Liters)"]), tc_style)],
                [Paragraph("FLUSHING WATER TANK", tc_left), Paragraph("1", tc_style), Paragraph(str(calculated_results[plot_name]["UGT Flushing (Liters)"]), tc_style)],
                [Paragraph("FIRE WATER TANK", tc_left), Paragraph("1", tc_style), Paragraph(str(calculated_results[plot_name]["Fire Tank (Liters)"]), tc_style)],
            ]
            t_u = Table(u_data, colWidths=[200, 120, 150])
            t_u.setStyle(TableStyle([('GRID', (0,0), (-1,-1), 0.5, colors.black), ('BACKGROUND', (0,0), (-1,0), colors.HexColor("#34495E"))]))
            story.append(t_u)
            story.append(PageBreak())

        # --- PAGE 5 & 6: UGT PLOT A & B ---
        make_ugt_page("Plot-A")
        make_ugt_page("Plot-B")

        # Function to generate STP Page
        def make_stp_page(plot_name):
            story.append(eng_header)
            story.append(Spacer(1, 5))
            story.append(Table([[Paragraph(f"<b>STP DETAILS - {plot_name}</b>", sec_style)]], colWidths=[554], style=[('BACKGROUND', (0,0), (-1,-1), colors.HexColor("#8E44AD"))]))
            story.append(Spacer(1, 5))
            s_data = [
                [Paragraph("<b>DESCRIPTION</b>", th_style), Paragraph("<b>CAPACITY</b>", th_style), Paragraph("<b>UNITS</b>", th_style)],
                [Paragraph("TOTAL WATER REQUIREMENT", tc_left), Paragraph(str(calculated_results[plot_name]["Total Water (LPD)"]), tc_style), Paragraph("LITERS/DAY", tc_style)],
                [Paragraph("SEWAGE GENERATION @90%", tc_left), Paragraph(str(calculated_results[plot_name]["Sewage Gen (LPD)"]), tc_style), Paragraph("LITERS/DAY", tc_style)],
                [Paragraph("SAY STP CAPACITY", tc_left), Paragraph(str(calculated_results[plot_name]["STP Capacity (KLD)"]), tc_style), Paragraph("KLD", tc_style)],
            ]
            t_s = Table(s_data, colWidths=[250, 120, 120])
            t_s.setStyle(TableStyle([('GRID', (0,0), (-1,-1), 0.5, colors.black), ('BACKGROUND', (0,0), (-1,0), colors.HexColor("#34495E"))]))
            story.append(t_s)
            story.append(PageBreak())

        # --- PAGE 7 & 8: STP PLOT A & B ---
        make_stp_page("Plot-A")
        make_stp_page("Plot-B")

        doc.build(story)
        messagebox.showinfo("Success", "8-Page Professional PDF Exported Successfully!")
    except Exception as e:
        messagebox.showerror("PDF Error", f"Failed to export PDF: {str(e)}")

def export_excel():
    file_path = filedialog.asksaveasfilename(defaultextension=".xlsx", filetypes=[("Excel files", "*.xlsx")])
    if not file_path:
        return
    try:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Summary"
        ws.append(["Project Summary"])
        ws.append(["Project Name", project_data.get('Project Name')])
        ws.append(["Total Water Demand (LPD)", calculated_results["Total"]["Total Water (LPD)"]])
        
        ws_a = wb.create_sheet("Plot-A")
        ws_a.append(["Plot-A Calculations"])
        for k, v in calculated_results["Plot-A"].items():
            ws_a.append([k, v])
            
        ws_b = wb.create_sheet("Plot-B")
        ws_b.append(["Plot-B Calculations"])
        for k, v in calculated_results["Plot-B"].items():
            ws_b.append([k, v])

        wb.save(file_path)
        messagebox.showinfo("Success", "Excel Exported Successfully with Plot Tabs!")
    except Exception as e:
        messagebox.showerror("Excel Error", str(e))

# ================= 1. PROJECT DETAILS PAGE =================
frame_proj = ctk.CTkFrame(container, fg_color="transparent")
frames["Project"] = frame_proj
frame_proj.grid(row=0, column=0, sticky="nsew")

ctk.CTkLabel(frame_proj, text="💧 WATER DEMAND REPORT GENERATOR", font=("Arial", 24, "bold")).pack(pady=10)
p_form = ctk.CTkFrame(frame_proj)
p_form.pack(fill="both", expand=True, padx=20, pady=10)

p_entries = {}
fields = [("Project Name", "PROPOSED RESIDENTIAL"), ("Client Name", "MR. ABC"), ("Project Location", "PUNE"), ("Engineer Name", "AKASH KHADE")]

for i, (lbl, default) in enumerate(fields):
    ctk.CTkLabel(p_form, text=lbl, font=("Arial", 14)).grid(row=i, column=0, padx=20, pady=10, sticky="w")
    ent = ctk.CTkEntry(p_form, width=350)
    ent.insert(0, default)
    ent.grid(row=i, column=1, padx=20, pady=10)
    p_entries[lbl] = ent

ctk.CTkLabel(p_form, text="Date", font=("Arial", 14)).grid(row=5, column=0, padx=20, pady=10, sticky="w")
date_entry = DateEntry(p_form, width=18, date_pattern="dd-mm-yyyy")
date_entry.grid(row=5, column=1, padx=20, pady=10, sticky="w")

def go_to_res():
    project_data.update({k: v.get() for k, v in p_entries.items()})
    project_data["Date"] = date_entry.get()
    show_frame("Residential")

ctk.CTkButton(frame_proj, text="Next -> Residential Details", command=go_to_res).pack(pady=15)

# ================= 2. RESIDENTIAL PAGE =================
frame_res = ctk.CTkFrame(container, fg_color="transparent")
frames["Residential"] = frame_res
frame_res.grid(row=0, column=0, sticky="nsew")

ctk.CTkLabel(frame_res, text="Residential Details (Plot A & B)", font=("Arial", 24, "bold")).pack(pady=10)
res_outer = ctk.CTkFrame(frame_res)
res_outer.pack(fill="both", expand=True, padx=20, pady=10)

headers = ["Plot", "Wing Name", "No. Of Flats", "Pop/Flat"]
for i, h in enumerate(headers):
    ctk.CTkLabel(res_outer, text=h, font=("Arial", 13, "bold")).grid(row=0, column=i, padx=5, pady=5)

res_rows = []
def add_res_row():
    r = len(res_rows) + 1
    plot_var = ctk.StringVar(value="Plot-A")
    plot_cb = ctk.CTkComboBox(res_outer, values=["Plot-A", "Plot-B"], variable=plot_var, width=90)
    w_ent = ctk.CTkEntry(res_outer, width=150)
    f_ent = ctk.CTkEntry(res_outer, width=100)
    p_ent = ctk.CTkEntry(res_outer, width=100)
    p_ent.insert(0, "5")
    
    plot_cb.grid(row=r, column=0, padx=5, pady=5)
    w_ent.grid(row=r, column=1, padx=5, pady=5)
    f_ent.grid(row=r, column=2, padx=5, pady=5)
    p_ent.grid(row=r, column=3, padx=5, pady=5)
    res_rows.append({"plot": plot_var, "wing": w_ent, "flats": f_ent, "pop": p_ent})

add_res_row()
add_res_row() # added two by default

def save_res_and_next():
    residential_data.clear()
    for row in res_rows:
        w = row["wing"].get()
        if w:
            f = int(row["flats"].get() or 0)
            p = int(row["pop"].get() or 0)
            residential_data.append({"Plot": row["plot"].get(), "Wing": w, "No. Of Flats": f, "Population": f*p})
    show_frame("Commercial")

res_btn = ctk.CTkFrame(frame_res, fg_color="transparent")
res_btn.pack(pady=10)
ctk.CTkButton(res_btn, text="+ Add Wing", command=add_res_row).grid(row=0, column=0, padx=10)
ctk.CTkButton(res_btn, text="Save & Next -> Commercial", command=save_res_and_next).grid(row=0, column=1, padx=10)

# ================= 3. COMMERCIAL PAGE =================
frame_com = ctk.CTkFrame(container, fg_color="transparent")
frames["Commercial"] = frame_com
frame_com.grid(row=0, column=0, sticky="nsew")

ctk.CTkLabel(frame_com, text="Commercial Details (Plot A & B)", font=("Arial", 24, "bold")).pack(pady=10)
com_outer = ctk.CTkFrame(frame_com)
com_outer.pack(fill="both", expand=True, padx=20, pady=10)

headers_c = ["Plot", "Type", "Area (sq.m)", "Pop (Auto)"]
for i, h in enumerate(headers_c):
    ctk.CTkLabel(com_outer, text=h, font=("Arial", 13, "bold")).grid(row=0, column=i, padx=5, pady=5)

com_rows = []
def add_com_row():
    r = len(com_rows) + 1
    plot_var = ctk.StringVar(value="Plot-A")
    plot_cb = ctk.CTkComboBox(com_outer, values=["Plot-A", "Plot-B"], variable=plot_var, width=90)
    t_var = ctk.StringVar(value="Shop")
    t_cb = ctk.CTkComboBox(com_outer, values=["Shop", "Office", "Restaurant"], variable=t_var, width=120)
    a_ent = ctk.CTkEntry(com_outer, width=100)
    p_var = ctk.StringVar(value="0")
    p_lbl = ctk.CTkLabel(com_outer, textvariable=p_var, width=100)
    
    def update_pop(*args):
        val = float(a_ent.get() or 0)
        p_var.set(str(int(val * 0.1))) # Assuming 0.1 factor for demo

    a_ent.bind("<KeyRelease>", update_pop)
    
    plot_cb.grid(row=r, column=0, padx=5, pady=5)
    t_cb.grid(row=r, column=1, padx=5, pady=5)
    a_ent.grid(row=r, column=2, padx=5, pady=5)
    p_lbl.grid(row=r, column=3, padx=5, pady=5)
    com_rows.append({"plot": plot_var, "type": t_var, "area": a_ent, "pop_var": p_var})

add_com_row()

def save_com_and_next():
    commercial_data.clear()
    for row in com_rows:
        a = row["area"].get()
        if a:
            commercial_data.append({"Plot": row["plot"].get(), "Commercial Type": row["type"].get(), "Area": float(a), "Auto Population": int(row["pop_var"].get())})
    show_frame("Other")

com_btn = ctk.CTkFrame(frame_com, fg_color="transparent")
com_btn.pack(pady=10)
ctk.CTkButton(com_btn, text="+ Add Commercial", command=add_com_row).grid(row=0, column=0, padx=10)
ctk.CTkButton(com_btn, text="Save & Next -> Other", command=save_com_and_next).grid(row=0, column=1, padx=10)

# ================= 4. OTHER DETAILS PAGE =================
frame_oth = ctk.CTkFrame(container, fg_color="transparent")
frames["Other"] = frame_oth
frame_oth.grid(row=0, column=0, sticky="nsew")

ctk.CTkLabel(frame_oth, text="Other / Fire Tank Details", font=("Arial", 24, "bold")).pack(pady=10)
oth_outer = ctk.CTkFrame(frame_oth)
oth_outer.pack(fill="both", expand=True, padx=20, pady=10)

other_fields = [("Fire Tank Plot-A", "litres"), ("Fire Tank Plot-B", "litres")]
oth_entries = {}
for i, (lbl, unit) in enumerate(other_fields, start=1):
    ctk.CTkLabel(oth_outer, text=lbl, font=("Arial", 14)).grid(row=i, column=0, padx=20, pady=10, sticky="w")
    ent = ctk.CTkEntry(oth_outer, width=200)
    ent.insert(0, "300000" if "A" in lbl else "230000")
    ent.grid(row=i, column=1, padx=20, pady=10)
    oth_entries[lbl] = ent

def calc_and_final():
    other_data.clear()
    for lbl, ent in oth_entries.items():
        other_data[lbl] = float(ent.get() or 0)
    perform_calculations()
    refresh_final()
    show_frame("Final")

ctk.CTkButton(frame_oth, text="Calculate & Generate Report", command=calc_and_final).pack(pady=20)

# ================= 5. FINAL REPORT PAGE =================
frame_final = ctk.CTkFrame(container, fg_color="transparent")
frames["Final"] = frame_final
frame_final.grid(row=0, column=0, sticky="nsew")

ctk.CTkLabel(frame_final, text="Calculations Ready!", font=("Arial", 24, "bold")).pack(pady=10)
res_lbl = ctk.CTkLabel(frame_final, text="", font=("Arial", 14), justify="center")
res_lbl.pack(pady=20)

def refresh_final():
    txt = f"Plot-A STP: {calculated_results['Plot-A']['STP Capacity (KLD)']} KLD\n"
    txt += f"Plot-B STP: {calculated_results['Plot-B']['STP Capacity (KLD)']} KLD\n\n"
    txt += f"Total Project Demand: {calculated_results['Total']['Total Water (LPD)']} LPD"
    res_lbl.configure(text=txt)

act_frame = ctk.CTkFrame(frame_final, fg_color="transparent")
act_frame.pack(pady=20)
ctk.CTkButton(act_frame, text="📄 Generate 8-Page PDF", fg_color="#C0392B", command=export_pdf).grid(row=0, column=0, padx=10)
ctk.CTkButton(act_frame, text="📊 Export Excel", fg_color="#27AE60", command=export_excel).grid(row=0, column=1, padx=10)
ctk.CTkButton(act_frame, text="<- Back to Edit", fg_color="gray", command=lambda: show_frame("Other")).grid(row=0, column=2, padx=10)

show_frame("Project")

if __name__ == "__main__":
    app.mainloop()

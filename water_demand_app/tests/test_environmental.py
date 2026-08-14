"""Point 2 & 3 — Solid waste and sewage generation tests."""

from __future__ import annotations

import os
import sys
import tempfile
import unittest

APP_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if APP_DIR not in sys.path:
    sys.path.insert(0, APP_DIR)

from config.nbc_2026 import PROJECT_TYPE_RESIDENTIAL
from config.page_visibility import visible_pages
from models.calculations import CalculationResults, PlotResults
from models.project import ProjectData
from models.residential import ResidentialWing
from openpyxl import load_workbook
from services.environmental_calculator import calculate_environmental, calculate_sewage_generation, calculate_solid_waste
from services.excel_exporter import TEMPLATE_PATH, export_excel
from services.result_tables import build_preview_table_sections, build_sewage_generation_table_sections, build_solid_waste_table_sections
from ui.app_state import AppState


def _reference_results(stp_kld: float = 370.0) -> CalculationResults:
    plot_a = PlotResults(
        plot="Plot-A",
        res_population=2835,
        com_population=0,
        stp_capacity_kld=stp_kld,
        dry_total_water_lpd=500000,
    )
    plot_b = PlotResults(plot="Plot-B")
    return CalculationResults(
        plots={"Plot-A": plot_a, "Plot-B": plot_b},
        total={
            "Total Residential Population": 2835,
            "Total Commercial Population": 0,
            "Total STP Capacity (KLD)": stp_kld,
            "Total Water (LPD)": 500000,
            "Total Population": 2835,
        },
    )


class TestSolidWasteCalculations(unittest.TestCase):
    def setUp(self) -> None:
        self.project = ProjectData(project_type=PROJECT_TYPE_RESIDENTIAL)
        self.results = _reference_results()
        self.env = calculate_environmental(self.results, self.project)

    def test_residential_waste_calculation(self) -> None:
        sw = self.env.solid_waste
        self.assertEqual(sw.res_population, 2835)
        self.assertAlmostEqual(sw.res_total_waste_kg_day, 1275.75, places=2)
        self.assertAlmostEqual(sw.res_wet_waste_kg_day, 765.45, places=2)
        self.assertAlmostEqual(sw.res_dry_waste_kg_day, 510.30, places=2)

    def test_commercial_waste_calculation(self) -> None:
        sw = self.env.solid_waste
        self.assertEqual(sw.comm_population, 0)
        self.assertAlmostEqual(sw.comm_total_waste_kg_day, 0.0, places=2)

    def test_total_waste_calculation(self) -> None:
        sw = self.env.solid_waste
        self.assertAlmostEqual(sw.total_waste_kg_day, 1275.75, places=2)
        self.assertAlmostEqual(sw.total_wet_waste_kg_day, 765.45, places=2)
        self.assertAlmostEqual(sw.total_dry_waste_kg_day, 510.30, places=2)

    def test_wet_waste_considered_and_owc(self) -> None:
        sw = self.env.solid_waste
        self.assertAlmostEqual(sw.wet_waste_considered_kg_day, 766.0, places=2)
        self.assertAlmostEqual(sw.garden_waste_considered_kg_day, 60.0, places=2)
        self.assertAlmostEqual(sw.stp_sludge_kg_day, 74.0, places=2)
        self.assertAlmostEqual(sw.owc_plant_capacity_kg_day, 900.0, places=2)

    def test_e_waste_calculation(self) -> None:
        sw = self.env.solid_waste
        self.assertAlmostEqual(sw.res_e_waste_kg_year, 2835.0, places=2)
        self.assertAlmostEqual(sw.total_e_waste_kg_year, 2835.0, places=2)
        self.assertEqual(sw.approx_area_sqm, "70-75")

    def test_solid_waste_table_sections_exist(self) -> None:
        sections = build_solid_waste_table_sections(self.project, self.env)
        titles = [s.title for s in sections]
        self.assertIn("Residential Waste Generated", titles)
        self.assertIn("Total E-Waste Generated", titles)

    def test_solid_waste_updates_when_population_changes(self) -> None:
        state = AppState()
        state.residential = [ResidentialWing(plot="Plot-A", wing="A", flats_2bhk=100)]
        state.project.project_type = PROJECT_TYPE_RESIDENTIAL
        state.auto_calculate()
        first = calculate_solid_waste(state.results, state.project).res_total_waste_kg_day
        state.residential[0].flats_2bhk = 200
        state.auto_calculate()
        second = calculate_solid_waste(state.results, state.project).res_total_waste_kg_day
        self.assertGreater(second, first)


class TestSewageGenerationCalculations(unittest.TestCase):
    def setUp(self) -> None:
        self.project = ProjectData(project_type=PROJECT_TYPE_RESIDENTIAL)
        self.results = _reference_results()
        self.env = calculate_environmental(self.results, self.project)

    def test_residential_sewage_calculation(self) -> None:
        sg = self.env.sewage
        self.assertAlmostEqual(sg.res_sewage_kld, 344.45, places=2)

    def test_commercial_sewage_calculation(self) -> None:
        sg = self.env.sewage
        self.assertAlmostEqual(sg.comm_sewage_kld, 0.0, places=2)

    def test_total_sewage_calculation(self) -> None:
        sg = self.env.sewage
        self.assertAlmostEqual(sg.total_sewage_kld, 344.45, places=2)

    def test_stp_capacity_integration(self) -> None:
        sg = self.env.sewage
        self.assertAlmostEqual(sg.proposed_stp_capacity_kld, 370.0, places=2)

    def test_stp_sludge_integration(self) -> None:
        sg = self.env.sewage
        self.assertAlmostEqual(sg.stp_sludge_kg_day, 74.0, places=2)
        self.assertEqual(sg.approx_area_sqm, "140-160")

    def test_sewage_table_sections_exist(self) -> None:
        sections = build_sewage_generation_table_sections(self.project, self.env)
        titles = [s.title for s in sections]
        self.assertIn("Residential Sewage Generated", titles)
        self.assertIn("Total Sewage Generated", titles)

    def test_sewage_updates_when_population_changes(self) -> None:
        state = AppState()
        state.residential = [ResidentialWing(plot="Plot-A", wing="A", flats_2bhk=50)]
        state.project.project_type = PROJECT_TYPE_RESIDENTIAL
        state.auto_calculate()
        first = calculate_sewage_generation(state.results, state.project).res_sewage_kld
        state.residential[0].flats_2bhk = 100
        state.auto_calculate()
        second = calculate_sewage_generation(state.results, state.project).res_sewage_kld
        self.assertGreater(second, first)


class TestEnvironmentalVisibilityAndExports(unittest.TestCase):
    def test_dynamic_tab_visibility(self) -> None:
        pages = visible_pages(PROJECT_TYPE_RESIDENTIAL)
        self.assertIn("Sewage", pages)
        self.assertIn("SolidWaste", pages)

    def test_preview_includes_environmental_sections(self) -> None:
        project = ProjectData(project_type=PROJECT_TYPE_RESIDENTIAL)
        results = _reference_results()
        env = calculate_environmental(results, project)
        sections = build_preview_table_sections(project, results, ("Plot-A", "Plot-B"), environmental=env)
        titles = [s.title for s in sections]
        self.assertTrue(any("Sewage" in t for t in titles))
        self.assertTrue(any("Solid Waste" in t for t in titles))

    def test_excel_includes_environmental_sheets(self) -> None:
        project = ProjectData(project_name="ENV TEST", project_type=PROJECT_TYPE_RESIDENTIAL)
        results = _reference_results()
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
            out_path = tmp.name
        try:
            export_excel(out_path, project, results)
            wb = load_workbook(out_path)
            self.assertIn("Sewage Generation", wb.sheetnames)
            self.assertIn("Solid Waste Generation", wb.sheetnames)
            if os.path.isfile(TEMPLATE_PATH):
                self.assertEqual(wb.sheetnames[0], "Cover")
        finally:
            os.unlink(out_path)

    def test_pdf_export_includes_environmental_builder(self) -> None:
        from services.pdf_exporter import PDFExporter

        project = ProjectData(project_name="ENV TEST", project_type=PROJECT_TYPE_RESIDENTIAL)
        results = _reference_results()
        env = calculate_environmental(results, project)
        exporter = PDFExporter(project, results)
        story = exporter._build_environmental_tables(
            "SEWAGE GENERATION",
            "#8E44AD",
            build_sewage_generation_table_sections(project, env),
        )
        self.assertGreater(len(story), 3)

    def test_pdf_export_completes(self) -> None:
        from services.pdf_exporter import export_pdf

        project = ProjectData(project_name="ENV TEST", project_type=PROJECT_TYPE_RESIDENTIAL)
        results = _reference_results()
        with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as tmp:
            out_path = tmp.name
        try:
            export_pdf(out_path, project, results)
            self.assertGreater(os.path.getsize(out_path), 5000)
        finally:
            os.unlink(out_path)


if __name__ == "__main__":
    unittest.main()

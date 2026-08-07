"""Tests for template-based Excel export."""

import os
import sys
import tempfile
import unittest

APP_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, APP_DIR)

from models.commercial import CommercialUnit
from models.other_details import OtherDetails
from models.project import ProjectData
from models.residential import ResidentialWing
from openpyxl import load_workbook
from services.calculator import WaterDemandCalculator
from services.excel_exporter import TEMPLATE_PATH, export_excel


class TestTemplateExcelExport(unittest.TestCase):
    def setUp(self) -> None:
        if not os.path.isfile(TEMPLATE_PATH):
            self.skipTest("WaterDemand_Template.xlsx not found")

    def test_template_export_preserves_sheets(self) -> None:
        residential = [
            ResidentialWing(plot="Plot-A", wing="WING - A", flats_2bhk=73, flats_3bhk=73, building_height_m=24),
        ]
        commercial = [
            CommercialUnit(plot="Plot-A", block="COMM-A", comm_type="Retail Shop", area_sqm=437),
        ]
        project = ProjectData(project_name="TEST PROJECT", client_name="CLIENT", engineer_name="Akash")
        calc = WaterDemandCalculator(residential, commercial, OtherDetails(), project)
        results = calc.calculate()

        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
            out_path = tmp.name
        try:
            export_excel(out_path, project, results)
            wb = load_workbook(out_path)
            self.assertEqual(
                wb.sheetnames,
                [
                    "Cover",
                    "Consolidated",
                    "Plot-A Demand",
                    "Plot-B Demand",
                    "Plot-A UGT-OHT",
                    "Plot-B UGT-OHT",
                    "Plot-A STP",
                    "Plot-B STP",
                    "Summary",
                ],
            )
            self.assertEqual(wb["Cover"]["B4"].value, "TEST PROJECT")
            self.assertEqual(wb["Cover"]["B8"].value, "Akash")
            self.assertGreater(wb["Plot-A Demand"]["E5"].value, 0)
            self.assertEqual(wb["Plot-A Demand"]["B16"].value, results.plots["Plot-A"].dry_total_water_lpd)
        finally:
            os.unlink(out_path)


if __name__ == "__main__":
    unittest.main()

"""Regression tests for OHT, STP Summary, and Preview table data."""

from __future__ import annotations

import os
import sys
import tempfile
import unittest

APP_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if APP_DIR not in sys.path:
    sys.path.insert(0, APP_DIR)

from config.nbc_2026 import PROJECT_TYPE_COMMERCIAL, PROJECT_TYPE_RESIDENTIAL
from config.page_visibility import visible_pages
from models.commercial import CommercialUnit
from models.other_details import OtherDetails
from models.project import ProjectData
from models.residential import ResidentialWing
from openpyxl import load_workbook
from services.calculator import WaterDemandCalculator
from services.excel_exporter import TEMPLATE_PATH, export_excel
from services.result_tables import (
    build_oht_table_sections,
    build_preview_table_sections,
    build_stp_table_sections,
    oht_totals_for_check,
    stp_rows_for_excel_check,
)
from ui.app_state import AppState


def _punavale_results():
    residential = [
        ResidentialWing(plot="Plot-A", wing="WING - A", flats=146, pop_per_flat=5),
        ResidentialWing(plot="Plot-A", wing="WING - B", flats=146, pop_per_flat=5),
        ResidentialWing(plot="Plot-B", wing="WING - C", flats=56, pop_per_flat=5),
    ]
    commercial = [
        CommercialUnit(
            plot="Plot-A",
            block="COMM-A",
            comm_type="Shop - Ground Floor",
            floor_label="Ground Floor (Shop)",
            area_sqm=437,
        ),
    ]
    other = OtherDetails(
        landscape_area={"Plot-A": 765.0, "Plot-B": 762.0},
        fire_tank={"Plot-A": 300000.0, "Plot-B": 230000.0},
    )
    project = ProjectData(project_name="Punavale Test", project_type=PROJECT_TYPE_RESIDENTIAL)
    return WaterDemandCalculator(residential, commercial, other, project).calculate(), project, other


def _section_titles(sections) -> list[str]:
    return [s.title for s in sections]


def _find_row(sections, description: str) -> tuple[str, str, str] | None:
    for section in sections:
        for row in section.rows:
            if row[0] == description:
                return row
    return None


class TestOHTResultTables(unittest.TestCase):
    def setUp(self) -> None:
        self.results, self.project, _ = _punavale_results()

    def test_oht_table_sections_exist(self) -> None:
        sections = build_oht_table_sections(self.results, ("Plot-A", "Plot-B"))
        self.assertGreaterEqual(len(sections), 1)
        self.assertTrue(any("OHT Details" in s.title for s in sections))

    def test_oht_values_from_calculation_engine(self) -> None:
        plot = self.results.plots["Plot-A"]
        self.assertTrue(plot.oht_rows)
        first = plot.oht_rows[0]
        sections = build_oht_table_sections(self.results, ("Plot-A",))
        row = _find_row(sections, f"{first['wing']} — Domestic OHT")
        self.assertIsNotNone(row)
        assert row is not None
        self.assertEqual(row[2], "KLD")
        self.assertIn(f"{first['domestic_kld']:.2f}", row[1].replace(",", ""))

    def test_oht_totals_match_engine(self) -> None:
        totals = oht_totals_for_check(self.results, "Plot-A")
        plot = self.results.plots["Plot-A"]
        expected_dom = sum(r["domestic_kld"] for r in plot.oht_rows)
        self.assertAlmostEqual(totals[0], expected_dom)


class TestSTPResultTables(unittest.TestCase):
    def setUp(self) -> None:
        self.results, self.project, _ = _punavale_results()

    def test_stp_table_sections_exist(self) -> None:
        sections = build_stp_table_sections(self.results, ("Plot-A", "Plot-B"))
        self.assertGreater(len(sections), 0)
        self.assertTrue(any("STP" in s.title for s in sections))

    def test_stp_values_from_calculation_engine(self) -> None:
        plot = self.results.plots["Plot-A"]
        self.assertTrue(plot.stp_sections)
        stp = plot.stp_sections[0]
        sections = build_stp_table_sections(self.results, ("Plot-A",))
        row = _find_row(sections, "Say STP Capacity")
        self.assertIsNotNone(row)
        assert row is not None
        self.assertEqual(row[2], "KLD")
        self.assertIn(f"{stp.say_stp_kld:.2f}", row[1])

    def test_stp_rows_match_excel_exporter_fields(self) -> None:
        rows = stp_rows_for_excel_check(self.results, "Plot-A")
        plot = self.results.plots["Plot-A"]
        stp = plot.stp_sections[0]
        sewage_row = next(r for r in rows if r[0] == "Sewage Generation @90%")
        self.assertEqual(sewage_row[1], f"{stp.sewage_lpd:,}")
        self.assertEqual(sewage_row[2], "LITERS/DAY")


class TestPreviewResultTables(unittest.TestCase):
    def setUp(self) -> None:
        self.results, self.project, self.other = _punavale_results()

    def test_preview_has_structured_sections(self) -> None:
        sections = build_preview_table_sections(
            self.project, self.results, ("Plot-A", "Plot-B"), other=self.other
        )
        titles = _section_titles(sections)
        self.assertIn("Project Summary", titles)
        self.assertIn("Water Demand Summary", titles)
        self.assertIn("STP Summary", titles)
        self.assertIn("OHT Summary", titles)

    def test_preview_values_match_calculation_results(self) -> None:
        sections = build_preview_table_sections(
            self.project, self.results, ("Plot-A", "Plot-B"), other=self.other
        )
        water_row = _find_row(sections, "Total Water Demand")
        self.assertIsNotNone(water_row)
        assert water_row is not None
        expected = f"{self.results.total['Total Water (LPD)']:,}"
        self.assertEqual(water_row[1], expected)
        self.assertEqual(water_row[2], "LPD")

    def test_preview_updates_when_inputs_change(self) -> None:
        state = AppState()
        state.residential = [ResidentialWing(plot="Plot-A", wing="A", flats_2bhk=10)]
        state.project.project_type = PROJECT_TYPE_RESIDENTIAL
        state.auto_calculate()
        first = build_preview_table_sections(
            state.project, state.results, ("Plot-A", "Plot-B"), other=state.other
        )
        first_water = _find_row(first, "Total Water Demand")
        state.residential[0].flats_2bhk = 30
        state.auto_calculate()
        second = build_preview_table_sections(
            state.project, state.results, ("Plot-A", "Plot-B"), other=state.other
        )
        second_water = _find_row(second, "Total Water Demand")
        self.assertIsNotNone(first_water)
        self.assertIsNotNone(second_water)
        assert first_water is not None and second_water is not None
        self.assertNotEqual(first_water[1], second_water[1])

    def test_irrelevant_sections_hidden_for_commercial_type(self) -> None:
        project = ProjectData(project_type=PROJECT_TYPE_COMMERCIAL)
        commercial = [
            CommercialUnit(plot="Plot-A", block="COMM-A", comm_type="Office", area_sqm=1000),
        ]
        results = WaterDemandCalculator([], commercial, OtherDetails(hvac_water={"Plot-A": 1000}), project).calculate()
        sections = build_preview_table_sections(project, results, ("Plot-A", "Plot-B"), other=OtherDetails())
        titles = _section_titles(sections)
        self.assertIn("STP Summary", titles)
        self.assertIn("OHT Summary", titles)
        pages = visible_pages(PROJECT_TYPE_COMMERCIAL)
        self.assertNotIn("Residential", pages)

    def test_residential_hides_hvac_in_other_section(self) -> None:
        other = OtherDetails(hvac_water={"Plot-A": 5000})
        sections = build_preview_table_sections(
            self.project, self.results, ("Plot-A",), other=other
        )
        titles = _section_titles(sections)
        if "Other Applicable Calculations" in titles:
            self.assertIsNone(_find_row(sections, "HVAC Makeup Water"))


class TestPreviewExcelConsistency(unittest.TestCase):
    def setUp(self) -> None:
        if not os.path.isfile(TEMPLATE_PATH):
            self.skipTest("WaterDemand_Template.xlsx not found")
        self.results, self.project, self.other = _punavale_results()

    def test_excel_stp_capacity_matches_preview(self) -> None:
        preview_sections = build_preview_table_sections(
            self.project, self.results, ("Plot-A", "Plot-B"), other=self.other
        )
        preview_stp = _find_row(preview_sections, "Proposed STP Capacity")
        self.assertIsNotNone(preview_stp)

        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
            out_path = tmp.name
        try:
            export_excel(out_path, self.project, self.results)
            wb = load_workbook(out_path)
            plot_a_stp = self.results.plots["Plot-A"].stp_capacity_kld
            total_stp = self.results.total["Total STP Capacity (KLD)"]
            excel_plot_a = None
            excel_total = None
            for row in range(3, 20):
                label = wb["Summary"].cell(row=row, column=1).value
                if label == "Plot-A STP Capacity (KLD)":
                    excel_plot_a = wb["Summary"].cell(row=row, column=2).value
                if label == "Total STP Capacity (KLD)":
                    excel_total = wb["Summary"].cell(row=row, column=2).value
            self.assertEqual(plot_a_stp, excel_plot_a)
            self.assertEqual(total_stp, excel_total)
            assert preview_stp is not None
            self.assertIn(f"{plot_a_stp:.2f}", preview_stp[1])
        finally:
            os.unlink(out_path)

    def test_preview_oht_totals_match_engine(self) -> None:
        sections = build_preview_table_sections(
            self.project, self.results, ("Plot-A",), other=self.other
        )
        dom_row = _find_row(sections, "Total Domestic OHT Capacity")
        self.assertIsNotNone(dom_row)
        totals = oht_totals_for_check(self.results, "Plot-A")
        assert dom_row is not None
        self.assertIn(f"{totals[0]:.2f}", dom_row[1].replace(",", ""))


class TestResultTableGUI(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        try:
            import customtkinter as ctk
            import importlib.util

            spec = importlib.util.spec_from_file_location(
                "water_main", os.path.join(os.path.dirname(APP_DIR), "main.py")
            )
            water_main = importlib.util.module_from_spec(spec)
            sys.modules["water_main"] = water_main
            assert spec.loader is not None
            spec.loader.exec_module(water_main)
            cls.app = water_main.WaterDemandApp()
            cls.app.withdraw()
            cls.app.update_idletasks()
            cls.app.update()
        except Exception as exc:
            raise unittest.SkipTest(f"GUI not available: {exc}") from exc

    @classmethod
    def tearDownClass(cls) -> None:
        if hasattr(cls, "app"):
            cls.app.destroy()

    def test_oht_result_table_widget_exists(self) -> None:
        self.assertTrue(hasattr(self.app, "oht_table"))

    def test_stp_result_table_widget_exists(self) -> None:
        self.assertTrue(hasattr(self.app, "stp_table"))

    def test_preview_result_table_widget_exists(self) -> None:
        self.assertTrue(hasattr(self.app, "preview_table"))

    def test_oht_page_refresh_populates_table(self) -> None:
        self.app.app_state.load_defaults()
        self.app.app_state.auto_calculate()
        self.app.show("OHT")
        self.app.update()
        sections = build_oht_table_sections(self.app.app_state.results, self.app._plots())
        self.assertGreater(len(sections), 0)


if __name__ == "__main__":
    unittest.main()

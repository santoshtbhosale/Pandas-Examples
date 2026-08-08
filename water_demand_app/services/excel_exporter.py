from __future__ import annotations

import os
import shutil
from typing import Any, Dict, Optional

from copy import copy

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from config.nbc_2026 import (
    BRAND_DARK_GRAY,
    BRAND_NAVY,
    BRAND_ORANGE,
    COMPANY_NAME,
    PLOT_MODE_SINGLE,
    active_plots,
    hvac_applicable,
)
from models.calculations import CalculationResults
from models.project import ProjectData


HEADER_FILL = PatternFill("solid", fgColor=BRAND_DARK_GRAY.replace("#", ""))
NAVY_FILL = PatternFill("solid", fgColor=BRAND_NAVY.replace("#", ""))
ORANGE_FILL = PatternFill("solid", fgColor=BRAND_ORANGE.replace("#", ""))
THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)
WHITE_BOLD = Font(color="FFFFFF", bold=True, size=10)
TITLE_FONT = Font(bold=True, size=14, color=BRAND_NAVY.replace("#", ""))
HEADER_FONT = Font(bold=True, size=11)


class ExcelExporter:
    def __init__(self, project: ProjectData, results: CalculationResults) -> None:
        self.project = project
        self.results = results
        self.wb = Workbook()

    def export(self, file_path: str) -> None:
        self._build_cover()
        self._build_consolidated()
        for plot in active_plots(self.project.plot_mode):
            self._build_plot_demand(plot)
            self._build_ugt_oht(plot)
            self._build_stp(plot)
        self._build_summary()
        if "Sheet" in self.wb.sheetnames:
            del self.wb["Sheet"]
        self.wb.save(file_path)

    def _style_header_row(self, ws, row: int, cols: int) -> None:
        for col in range(1, cols + 1):
            cell = ws.cell(row=row, column=col)
            cell.fill = HEADER_FILL
            cell.font = WHITE_BOLD
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = THIN_BORDER

    def _style_data_area(self, ws, start_row: int, end_row: int, cols: int) -> None:
        for r in range(start_row, end_row + 1):
            for c in range(1, cols + 1):
                cell = ws.cell(row=r, column=c)
                cell.border = THIN_BORDER
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    def _auto_width(self, ws) -> None:
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                if cell.value:
                    max_len = max(max_len, len(str(cell.value)))
            ws.column_dimensions[col_letter].width = min(max_len + 3, 45)

    def _build_cover(self) -> None:
        ws = self.wb.active
        ws.title = "Cover"
        ws["A1"] = COMPANY_NAME
        ws["A1"].font = TITLE_FONT
        ws.merge_cells("A1:F1")
        ws["A3"] = "TITLE"
        ws["B3"] = "WATER DEMAND"
        ws["A4"] = "PROJECT NAME"
        ws["B4"] = self.project.project_name
        ws["A5"] = "CLIENT NAME"
        ws["B5"] = self.project.client_name
        ws["A6"] = "PROJECT LOCATION"
        ws["B6"] = self.project.project_location
        ws["A7"] = "PROJECT NO."
        ws["B7"] = self.project.project_no
        ws["A8"] = "ENGINEER"
        ws["B8"] = self.project.engineer_name
        ws["A9"] = "DATE"
        ws["B9"] = self.project.date
        rev = self.project.revision
        headers = ["DATE", "REV. NO.", "DESCRIPTION", "PRPD. BY", "CHKD. BY", "APPRD. BY"]
        for i, h in enumerate(headers, 1):
            ws.cell(row=11, column=i, value=h)
        self._style_header_row(ws, 11, 6)
        ws.cell(row=12, column=1, value=rev.date or self.project.date)
        ws.cell(row=12, column=2, value=rev.revision_no)
        ws.cell(row=12, column=3, value=rev.description)
        ws.cell(row=12, column=4, value=rev.prepared_by)
        ws.cell(row=12, column=5, value=rev.checked_by)
        ws.cell(row=12, column=6, value=rev.approved_by)
        self._style_data_area(ws, 12, 12, 6)
        self._auto_width(ws)

    def _build_consolidated(self) -> None:
        ws = self.wb.create_sheet("Consolidated")
        ws["A1"] = "CONSOLIDATED STATEMENT"
        ws["A1"].font = HEADER_FONT
        ws.merge_cells("A1:J1")
        ws["A1"].fill = ORANGE_FILL
        headers = [
            "SR.NO", "DESCRIPTION", "PLOT-A RES", "PLOT-A COMM", "PLOT-A SUB",
            "PLOT-B RES", "PLOT-B COMM", "PLOT-B SUB", "TOTAL", "UNITS",
        ]
        for i, h in enumerate(headers, 1):
            ws.cell(row=3, column=i, value=h)
        self._style_header_row(ws, 3, 10)
        pa = self.results.plots["Plot-A"]
        pb = self.results.plots["Plot-B"]
        tot = self.results.total
        rows = [
            ("Number Of Building", pa.num_buildings_res, pa.num_buildings_com, pa.num_buildings_res + pa.num_buildings_com,
             pb.num_buildings_res, pb.num_buildings_com, pb.num_buildings_res + pb.num_buildings_com,
             pa.num_buildings_res + pa.num_buildings_com + pb.num_buildings_res + pb.num_buildings_com, "NO.S"),
            ("Total Number Of Flats", pa.total_flats, 0, pa.total_flats, pb.total_flats, 0, pb.total_flats, tot.get("Total Flats", 0), "NO.S"),
            ("Total Population", pa.res_population, pa.com_population, pa.total_population,
             pb.res_population, pb.com_population, pb.total_population, tot.get("Total Population", 0), "NO.S"),
            ("Total Water Demand (Dry)", pa.res_total_lpd + pa.com_total_lpd, 0, pa.dry_total_water_lpd,
             pb.res_total_lpd + pb.com_total_lpd, 0, pb.dry_total_water_lpd, tot.get("Total Water (LPD)", 0), "LPD"),
            ("STP Capacity", pa.stp_capacity_kld, 0, pa.stp_capacity_kld,
             pb.stp_capacity_kld, 0, pb.stp_capacity_kld, tot.get("Total STP Capacity (KLD)", 0), "KLD"),
        ]
        for ridx, row in enumerate(rows, 4):
            ws.cell(row=ridx, column=1, value=ridx - 3)
            for cidx, val in enumerate(row, 2):
                ws.cell(row=ridx, column=cidx, value=val)
        self._style_data_area(ws, 4, 4 + len(rows) - 1, 10)
        self._auto_width(ws)

    def _build_plot_demand(self, plot_name: str) -> None:
        plot = self.results.plots[plot_name]
        ws = self.wb.create_sheet(f"{plot_name} Demand")
        ws["A1"] = f"WATER DEMAND - {plot_name}"
        ws["A1"].font = HEADER_FONT
        ws.merge_cells("A1:H1")

        ws["A3"] = "RESIDENTIAL"
        ws["A3"].font = Font(bold=True)
        res_headers = ["SR.NO", "WING", "FLATS", "POP/FLAT", "POPULATION", "DOMESTIC", "FLUSHING", "KITCHEN", "TOTAL"]
        for i, h in enumerate(res_headers, 1):
            ws.cell(row=4, column=i, value=h)
        self._style_header_row(ws, 4, 9)
        row = 5
        for idx, w in enumerate(plot.residential_wings, 1):
            ws.cell(row=row, column=1, value=idx)
            ws.cell(row=row, column=2, value=w.wing)
            ws.cell(row=row, column=3, value=w.flats)
            ws.cell(row=row, column=4, value=w.pop_per_flat)
            ws.cell(row=row, column=5, value=w.population)
            ws.cell(row=row, column=6, value=w.domestic_lpd)
            ws.cell(row=row, column=7, value=w.flushing_lpd)
            ws.cell(row=row, column=8, value=w.kitchen_water_lpd)
            ws.cell(row=row, column=9, value=w.total_lpd)
            row += 1
        ws.cell(row=row, column=2, value="SUB-TOTAL")
        ws.cell(row=row, column=5, value=plot.res_population)
        ws.cell(row=row, column=6, value=plot.res_domestic_lpd)
        ws.cell(row=row, column=7, value=plot.res_flushing_lpd)
        ws.cell(row=row, column=8, value=plot.kitchen_water_lpd)
        ws.cell(row=row, column=9, value=plot.res_total_lpd)
        self._style_data_area(ws, 5, row, 9)
        row += 2

        ws.cell(row=row, column=1, value="COMMERCIAL")
        ws.cell(row=row, column=1).font = Font(bold=True)
        row += 1
        com_headers = ["SR.NO", "BLOCK", "TYPE/FLOOR", "AREA", "DENSITY", "POPULATION", "DOMESTIC", "FLUSHING", "TOTAL"]
        for i, h in enumerate(com_headers, 1):
            ws.cell(row=row, column=i, value=h)
        self._style_header_row(ws, row, 9)
        row += 1
        start_com = row
        for idx, u in enumerate(plot.commercial_units, 1):
            ws.cell(row=row, column=1, value=idx)
            ws.cell(row=row, column=2, value=u.block)
            ws.cell(row=row, column=3, value=u.floor_label or u.comm_type)
            ws.cell(row=row, column=4, value=u.area_sqm)
            ws.cell(row=row, column=5, value=u.density)
            ws.cell(row=row, column=6, value=u.population)
            ws.cell(row=row, column=7, value=u.domestic_lpd)
            ws.cell(row=row, column=8, value=u.flushing_lpd)
            ws.cell(row=row, column=9, value=u.total_lpd)
            row += 1
        self._style_data_area(ws, start_com, row - 1, 9)
        row += 1
        ws.cell(row=row, column=1, value="OTHER MEASURES")
        ws.cell(row=row, column=1).font = Font(bold=True)
        row += 1
        ws.cell(row=row, column=1, value="Landscape (NBC-2026)")
        ws.cell(row=row, column=2, value=plot.landscape_dry_lpd)
        ws.cell(row=row + 1, column=1, value="Swimming Pool")
        ws.cell(row=row + 1, column=2, value=plot.swimming_pool_lpd if plot.swimming_pool_lpd else "NA")
        next_row = row + 2
        if hvac_applicable(self.project.project_type):
            ws.cell(row=next_row, column=1, value="HVAC")
            ws.cell(row=next_row, column=2, value=plot.hvac_lpd)
            next_row += 1
        ws.cell(row=next_row, column=1, value="Kitchen Water")
        ws.cell(row=next_row, column=2, value=plot.kitchen_water_lpd)
        ws.cell(row=next_row + 1, column=1, value="GRAND TOTAL")
        ws.cell(row=next_row + 1, column=1).font = Font(bold=True)
        ws.cell(row=next_row + 1, column=2, value=plot.dry_total_water_lpd)
        self._auto_width(ws)

    def _build_ugt_oht(self, plot_name: str) -> None:
        plot = self.results.plots[plot_name]
        ws = self.wb.create_sheet(f"{plot_name} UGT-OHT")
        ws["A1"] = f"UGT & OHT DETAILS - {plot_name}"
        ws["A1"].font = HEADER_FONT
        ws.merge_cells("A1:F1")
        headers = ["SR.NO", "DESCRIPTION", "WATER REQ (LPD)", "STORAGE DAYS", "TOTAL STORAGE (L)", "TOTAL (KLD)"]
        for i, h in enumerate(headers, 1):
            ws.cell(row=3, column=i, value=h)
        self._style_header_row(ws, 3, 6)
        row = 4
        for idx, sec in enumerate(plot.ugt_sections, 1):
            ws.cell(row=row, column=1, value=idx)
            ws.cell(row=row, column=2, value=sec.description)
            ws.cell(row=row, column=3, value=sec.water_requirement_lpd)
            ws.cell(row=row, column=4, value=sec.storage_days)
            ws.cell(row=row, column=5, value=sec.total_storage_liters)
            ws.cell(row=row, column=6, value=sec.total_storage_kld)
            row += 1
        self._style_data_area(ws, 4, row - 1, 6)
        row += 2
        ws.cell(row=row, column=1, value="OHT DETAILS")
        ws.cell(row=row, column=1).font = Font(bold=True)
        row += 1
        oht_headers = ["SR.NO", "WING", "DOMESTIC KLD", "FLUSHING KLD", "FIRE BREAK KLD", "FIRE OHT KLD"]
        for i, h in enumerate(oht_headers, 1):
            ws.cell(row=row, column=i, value=h)
        self._style_header_row(ws, row, 6)
        row += 1
        start = row
        for idx, oht in enumerate(plot.oht_rows, 1):
            ws.cell(row=row, column=1, value=idx)
            ws.cell(row=row, column=2, value=oht["wing"])
            ws.cell(row=row, column=3, value=oht["domestic_kld"])
            ws.cell(row=row, column=4, value=oht["flushing_kld"])
            ws.cell(row=row, column=5, value=oht["fire_break_kld"])
            ws.cell(row=row, column=6, value=oht["fire_oht_kld"])
            row += 1
        self._style_data_area(ws, start, row - 1, 6)
        self._auto_width(ws)

    def _build_stp(self, plot_name: str) -> None:
        plot = self.results.plots[plot_name]
        ws = self.wb.create_sheet(f"{plot_name} STP")
        ws["A1"] = f"STP DETAILS - {plot_name}"
        ws["A1"].font = HEADER_FONT
        ws.merge_cells("A1:D1")
        row = 3
        for stp in plot.stp_sections:
            ws.cell(row=row, column=1, value=f"STP FOR {stp.scope}")
            ws.cell(row=row, column=1).font = Font(bold=True)
            row += 1
            headers = ["SR.NO", "DESCRIPTION", "CAPACITY", "UNITS"]
            for i, h in enumerate(headers, 1):
                ws.cell(row=row, column=i, value=h)
            self._style_header_row(ws, row, 4)
            row += 1
            data = [
                ("Total Water Requirement", stp.total_water_lpd, "LITERS/DAY"),
                ("Sewage Generation @90%", stp.sewage_lpd, "LITERS/DAY"),
                ("Capacity of Sewage Generation", stp.sewage_kld, "KLD"),
                ("Say STP Capacity", stp.say_stp_kld, "KLD"),
                ("Treated Water After Filtration", stp.treated_water_lpd, "LITERS/DAY"),
                ("Reuse Water for Flushing", stp.reuse_flushing_lpd, "LITERS/DAY"),
                ("Reuse Water for Landscape", stp.reuse_landscape_lpd, "LITERS/DAY"),
                ("Reuse Water for HVAC", stp.reuse_hvac_lpd, "LITERS/DAY"),
                ("Excess Treated Water", stp.excess_treated_lpd, "LITERS/DAY"),
            ]
            start = row
            for idx, (desc, cap, unit) in enumerate(data, 1):
                ws.cell(row=row, column=1, value=idx)
                ws.cell(row=row, column=2, value=desc)
                ws.cell(row=row, column=3, value=cap)
                ws.cell(row=row, column=4, value=unit)
                row += 1
            self._style_data_area(ws, start, row - 1, 4)
            row += 2
        self._auto_width(ws)

    def _build_summary(self) -> None:
        ws = self.wb.create_sheet("Summary")
        ws["A1"] = "PROJECT SUMMARY"
        ws["A1"].font = TITLE_FONT
        tot = self.results.total
        pa = self.results.plots["Plot-A"]
        pb = self.results.plots["Plot-B"]
        summary = [
            ("Project Name", self.project.project_name),
            ("Client Name", self.project.client_name),
            ("Project Location", self.project.project_location),
            ("Engineer", self.project.engineer_name),
            ("Date", self.project.date),
            ("", ""),
            ("Plot-A STP Capacity (KLD)", pa.stp_capacity_kld),
            ("Plot-B STP Capacity (KLD)", pb.stp_capacity_kld),
            ("Total Water Demand (LPD)", tot.get("Total Water (LPD)", 0)),
            ("Total STP Capacity (KLD)", tot.get("Total STP Capacity (KLD)", 0)),
            ("Total Population", tot.get("Total Population", 0)),
            ("Total Flats", tot.get("Total Flats", 0)),
        ]
        for idx, (k, v) in enumerate(summary, 3):
            ws.cell(row=idx, column=1, value=k)
            ws.cell(row=idx, column=2, value=v)
            if k:
                ws.cell(row=idx, column=1).font = Font(bold=True)
        self._auto_width(ws)


TEMPLATE_PATH = os.path.normpath(
    os.path.join(os.path.dirname(os.path.dirname(__file__)), "..", "templates", "WaterDemand_Template.xlsx")
)


class TemplateExcelExporter:
    """Populate the master Excel template — formatting preserved, data cells only."""

    def __init__(self, project: ProjectData, results: CalculationResults) -> None:
        self.project = project
        self.results = results

    def export(self, file_path: str) -> None:
        shutil.copy2(TEMPLATE_PATH, file_path)
        wb = load_workbook(file_path)
        self._populate_cover(wb)
        self._populate_consolidated(wb)
        for plot in ("Plot-A", "Plot-B"):
            if plot in active_plots(self.project.plot_mode):
                self._populate_plot_demand(wb, plot)
                self._populate_ugt_oht(wb, plot)
                self._populate_stp(wb, plot)
            else:
                self._clear_plot_sheets(wb, plot)
        self._populate_summary(wb)
        wb.save(file_path)

    @staticmethod
    def _set(ws, cell: str, value: Any) -> None:
        target = ws[cell]
        if type(target).__name__ == "MergedCell":
            for merged in ws.merged_cells.ranges:
                if target.coordinate in merged:
                    ws.cell(row=merged.min_row, column=merged.min_col).value = value
                    return
        target.value = value

    @staticmethod
    def _find_row(
        ws,
        needle: str,
        column: int = 1,
        start: int = 1,
        exact: bool = False,
    ) -> Optional[int]:
        needle_l = needle.lower().strip()
        for row in range(start, ws.max_row + 1):
            val = ws.cell(row=row, column=column).value
            if val is None:
                continue
            text = str(val).lower().strip()
            if exact:
                if text == needle_l:
                    return row
            elif needle_l in text:
                return row
        return None

    @staticmethod
    def _clear_range(ws, start_row: int, end_row: int, start_col: int, end_col: int) -> None:
        for row in range(start_row, end_row + 1):
            for col in range(start_col, end_col + 1):
                cell = ws.cell(row=row, column=col)
                if cell.value is not None and not str(cell.value).startswith("="):
                    cell.value = None

    def _populate_cover(self, wb) -> None:
        if "Cover" not in wb.sheetnames:
            return
        ws = wb["Cover"]
        rev = self.project.revision
        fields = {
            "B4": self.project.project_name,
            "B5": self.project.client_name,
            "B6": self.project.project_location,
            "B7": self.project.project_no,
            "B8": self.project.engineer_name,
            "B9": self.project.date,
            "A12": rev.date or self.project.date,
            "B12": rev.revision_no,
            "C12": rev.description,
            "D12": rev.prepared_by,
            "E12": rev.checked_by,
            "F12": rev.approved_by,
        }
        for cell, val in fields.items():
            self._set(ws, cell, val)

    def _populate_consolidated(self, wb) -> None:
        if "Consolidated" not in wb.sheetnames:
            return
        ws = wb["Consolidated"]
        pa = self.results.plots["Plot-A"]
        pb = self.results.plots["Plot-B"]
        tot = self.results.total
        single = self.project.plot_mode == PLOT_MODE_SINGLE
        rows = [
            (
                4,
                pa.num_buildings_res,
                pa.num_buildings_com,
                pa.num_buildings_res + pa.num_buildings_com,
                0 if single else pb.num_buildings_res,
                0 if single else pb.num_buildings_com,
                0 if single else pb.num_buildings_res + pb.num_buildings_com,
                pa.num_buildings_res + pa.num_buildings_com
                + (0 if single else pb.num_buildings_res + pb.num_buildings_com),
            ),
            (
                5,
                pa.total_flats,
                0,
                pa.total_flats,
                0 if single else pb.total_flats,
                0,
                0 if single else pb.total_flats,
                tot.get("Total Flats", pa.total_flats),
            ),
            (
                6,
                pa.res_population,
                pa.com_population,
                pa.total_population,
                0 if single else pb.res_population,
                0 if single else pb.com_population,
                0 if single else pb.total_population,
                tot.get("Total Population", pa.total_population),
            ),
            (
                7,
                pa.res_total_lpd + pa.com_total_lpd,
                0,
                pa.dry_total_water_lpd,
                0 if single else pb.res_total_lpd + pb.com_total_lpd,
                0,
                0 if single else pb.dry_total_water_lpd,
                tot.get("Total Water (LPD)", pa.dry_total_water_lpd),
            ),
            (
                8,
                pa.stp_capacity_kld,
                0,
                pa.stp_capacity_kld,
                0 if single else pb.stp_capacity_kld,
                0,
                0 if single else pb.stp_capacity_kld,
                tot.get("Total STP Capacity (KLD)", pa.stp_capacity_kld),
            ),
        ]
        for row, c, d, e, f, g, h, i in rows:
            for col, val in zip("CDEFGHI", (c, d, e, f, g, h, i)):
                self._set(ws, f"{col}{row}", val)

    @staticmethod
    def _copy_row_style(ws, src_row: int, dest_row: int, max_col: int = 9) -> None:
        for col in range(1, max_col + 1):
            src = ws.cell(row=src_row, column=col)
            dst = ws.cell(row=dest_row, column=col)
            dst.value = None
            if src.has_style:
                dst._style = copy(src._style)
            dst.number_format = src.number_format

    @staticmethod
    def _insert_rows_before(ws, before_row: int, count: int, style_row: int, max_col: int = 9) -> None:
        if count <= 0:
            return
        ws.insert_rows(before_row, count)
        TemplateExcelExporter._copy_row_style(ws, style_row, style_row, max_col)
        for offset in range(count):
            TemplateExcelExporter._copy_row_style(ws, style_row, before_row + offset, max_col)

    def _populate_plot_demand(self, wb, plot_name: str) -> None:
        sheet_name = f"{plot_name} Demand"
        if sheet_name not in wb.sheetnames:
            return
        ws = wb[sheet_name]
        plot = self.results.plots[plot_name]

        res_hdr = self._find_row(ws, "SR.NO", start=1)
        sub_row = self._find_row(ws, "SUB-TOTAL", column=2, exact=True)
        if res_hdr and sub_row:
            data_start = res_hdr + 1
            wings = plot.residential_wings
            available = max(0, sub_row - data_start)
            if len(wings) > available:
                self._insert_rows_before(ws, sub_row, len(wings) - available, data_start, 8)
                sub_row = self._find_row(ws, "SUB-TOTAL", column=2, exact=True)
            self._clear_range(ws, data_start, sub_row - 1, 1, 8)
            row = data_start
            for idx, wing in enumerate(wings, 1):
                pop_per = wing.pop_per_flat
                if wing.flats > 0 and wing.population > 0:
                    pop_per = round(wing.population / wing.flats)
                self._set(ws, f"A{row}", idx)
                self._set(ws, f"B{row}", wing.wing)
                self._set(ws, f"C{row}", wing.flats)
                self._set(ws, f"D{row}", pop_per)
                self._set(ws, f"E{row}", wing.population)
                self._set(ws, f"F{row}", wing.domestic_lpd)
                self._set(ws, f"G{row}", wing.flushing_lpd)
                self._set(ws, f"H{row}", wing.total_lpd)
                row += 1
            self._set(ws, f"E{sub_row}", plot.res_population)
            self._set(ws, f"F{sub_row}", plot.res_domestic_lpd)
            self._set(ws, f"G{sub_row}", plot.res_flushing_lpd)
            self._set(ws, f"H{sub_row}", plot.res_total_lpd)

        comm_row = self._find_row(ws, "COMMERCIAL", exact=True)
        other_row = self._find_row(ws, "OTHER MEASURES", exact=True)
        if comm_row and other_row:
            hdr = comm_row + 1
            data_start = hdr + 1
            units = plot.commercial_units
            gap_before_other = 1 if other_row - data_start > 0 else 0
            available = max(0, other_row - data_start - gap_before_other)
            if len(units) > available:
                self._insert_rows_before(ws, other_row, len(units) - available, data_start, 9)
                other_row = self._find_row(ws, "OTHER MEASURES", exact=True)
            self._clear_range(ws, data_start, other_row - 2, 1, 9)
            row = data_start
            for idx, unit in enumerate(units, 1):
                label = unit.floor_label or unit.comm_type
                self._set(ws, f"A{row}", idx)
                self._set(ws, f"B{row}", unit.block)
                self._set(ws, f"C{row}", label)
                self._set(ws, f"D{row}", unit.area_sqm)
                self._set(ws, f"E{row}", unit.density)
                self._set(ws, f"F{row}", unit.population)
                self._set(ws, f"G{row}", unit.domestic_lpd)
                self._set(ws, f"H{row}", unit.flushing_lpd)
                self._set(ws, f"I{row}", unit.total_lpd)
                row += 1

        other_row = self._find_row(ws, "OTHER MEASURES", exact=True)
        if other_row:
            for label, value in (
                ("Landscape", plot.landscape_dry_lpd),
                ("Swimming Pool", plot.swimming_pool_lpd if plot.swimming_pool_lpd else "NA"),
                ("HVAC", plot.hvac_lpd),
                ("Kitchen", plot.kitchen_water_lpd),
            ):
                row = self._find_row(ws, label, start=other_row)
                if row:
                    self._set(ws, f"B{row}", value)
            grand = self._find_row(ws, "GRAND TOTAL", start=other_row, exact=True)
            if grand:
                self._set(ws, f"B{grand}", plot.dry_total_water_lpd)

    def _populate_ugt_oht(self, wb, plot_name: str) -> None:
        sheet_name = f"{plot_name} UGT-OHT"
        if sheet_name not in wb.sheetnames:
            return
        ws = wb[sheet_name]
        plot = self.results.plots[plot_name]
        oht_hdr = self._find_row(ws, "OHT DETAILS", start=3, exact=True)
        ugt_end = (oht_hdr - 2) if oht_hdr else 20
        sections = plot.ugt_sections
        available = max(0, ugt_end - 4 + 1)
        if len(sections) > available and oht_hdr:
            self._insert_rows_before(ws, oht_hdr, len(sections) - available, 4, 6)
            oht_hdr = self._find_row(ws, "OHT DETAILS", start=3, exact=True)
            ugt_end = oht_hdr - 2
        self._clear_range(ws, 4, ugt_end, 1, 6)
        row = 4
        for idx, sec in enumerate(plot.ugt_sections, 1):
            self._set(ws, f"A{row}", idx)
            self._set(ws, f"B{row}", sec.description)
            self._set(ws, f"C{row}", sec.water_requirement_lpd)
            self._set(ws, f"D{row}", sec.storage_days)
            self._set(ws, f"E{row}", sec.total_storage_liters)
            self._set(ws, f"F{row}", sec.total_storage_kld)
            row += 1

        if oht_hdr:
            hdr = oht_hdr + 1
            oht_rows = plot.oht_rows
            available = 20
            if len(oht_rows) > available:
                self._insert_rows_before(ws, hdr + available + 1, len(oht_rows) - available, hdr + 1, 6)
            self._clear_range(ws, hdr + 1, hdr + max(len(oht_rows), 20), 1, 6)
            row = hdr + 1
            for idx, oht in enumerate(plot.oht_rows, 1):
                self._set(ws, f"A{row}", idx)
                self._set(ws, f"B{row}", oht["wing"])
                self._set(ws, f"C{row}", oht["domestic_kld"])
                self._set(ws, f"D{row}", oht["flushing_kld"])
                self._set(ws, f"E{row}", oht["fire_break_kld"])
                self._set(ws, f"F{row}", oht["fire_oht_kld"])
                row += 1

    def _populate_stp(self, wb, plot_name: str) -> None:
        sheet_name = f"{plot_name} STP"
        if sheet_name not in wb.sheetnames:
            return
        ws = wb[sheet_name]
        plot = self.results.plots[plot_name]
        stp_labels = [
            "Total Water Requirement",
            "Sewage Generation @90%",
            "Capacity of Sewage Generation",
            "Say STP Capacity",
            "Treated Water After Filtration",
            "Reuse Water for Flushing",
            "Reuse Water for Landscape",
            "Reuse Water for HVAC",
            "Excess Treated Water",
        ]
        stp_data = []
        for stp in plot.stp_sections:
            stp_data.append(
                (
                    f"STP FOR {stp.scope}",
                    [
                        (stp.total_water_lpd, "LITERS/DAY"),
                        (stp.sewage_lpd, "LITERS/DAY"),
                        (stp.sewage_kld, "KLD"),
                        (stp.say_stp_kld, "KLD"),
                        (stp.treated_water_lpd, "LITERS/DAY"),
                        (stp.reuse_flushing_lpd, "LITERS/DAY"),
                        (stp.reuse_landscape_lpd, "LITERS/DAY"),
                        (stp.reuse_hvac_lpd, "LITERS/DAY"),
                        (stp.excess_treated_lpd, "LITERS/DAY"),
                    ],
                )
            )

        row = 3
        for title, rows in stp_data:
            title_row = self._find_row(ws, title, start=row)
            if not title_row:
                ws.cell(row=row, column=1, value=title)
                title_row = row
            hdr = title_row + 1
            data_start = hdr + 1
            self._clear_range(ws, data_start, data_start + 8, 1, 4)
            for idx, (desc, (cap, unit)) in enumerate(zip(stp_labels, rows), 1):
                r = data_start + idx - 1
                self._set(ws, f"A{r}", idx)
                self._set(ws, f"B{r}", desc)
                self._set(ws, f"C{r}", cap)
                self._set(ws, f"D{r}", unit)
            row = data_start + 12

    def _clear_plot_sheets(self, wb, plot_name: str) -> None:
        for suffix in (" Demand", " UGT-OHT", " STP"):
            name = f"{plot_name}{suffix}"
            if name not in wb.sheetnames:
                continue
            ws = wb[name]
            self._clear_range(ws, 5, ws.max_row, 1, 12)

    def _populate_summary(self, wb) -> None:
        if "Summary" not in wb.sheetnames:
            return
        ws = wb["Summary"]
        tot = self.results.total
        pa = self.results.plots["Plot-A"]
        pb = self.results.plots["Plot-B"]
        single = self.project.plot_mode == PLOT_MODE_SINGLE
        fields = {
            "B3": self.project.project_name,
            "B4": self.project.client_name,
            "B5": self.project.project_location,
            "B6": self.project.engineer_name,
            "B7": self.project.date,
            "B9": pa.stp_capacity_kld,
            "B10": 0 if single else pb.stp_capacity_kld,
            "B11": tot.get("Total Water (LPD)", pa.dry_total_water_lpd),
            "B12": tot.get("Total STP Capacity (KLD)", pa.stp_capacity_kld),
            "B13": tot.get("Total Population", pa.total_population),
            "B14": tot.get("Total Flats", pa.total_flats),
        }
        for cell, val in fields.items():
            self._set(ws, cell, val)


def export_excel(file_path: str, project: ProjectData, results: CalculationResults) -> None:
    if os.path.isfile(TEMPLATE_PATH):
        TemplateExcelExporter(project, results).export(file_path)
    else:
        ExcelExporter(project, results).export(file_path)

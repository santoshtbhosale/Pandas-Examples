#!/usr/bin/env python3
"""Build single-file main.py from water_demand_app modules."""
import ast
import os

BASE = "water_demand_app"
OUT = "main.py"

HEADER = '''#!/usr/bin/env python3
"""
American Edge Engineers - Water Demand Report Generator
Single-file production application with NBC-2026 calculations.
Version 2.0.0
"""
from __future__ import annotations

import json
import math
import os
import re
import secrets
import hashlib
import shutil
import sqlite3
import uuid
from collections import defaultdict
from dataclasses import dataclass, field, asdict
from datetime import datetime
from tkinter import filedialog, messagebox
from typing import Any, Callable, Dict, List, Optional, Tuple, TypeVar

import customtkinter as ctk
from tkcalendar import DateEntry
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Image, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

APP_DIR = os.path.dirname(os.path.abspath(__file__))
DB_PATH = os.path.join(APP_DIR, "water_demand.db")
LOGO_PATH = os.path.join(APP_DIR, "logo.png")
JSON_SCHEMA_VERSION = "1.0"

'''


def remove_imports(source: str) -> str:
    tree = ast.parse(source)
    lines = source.splitlines(keepends=True)
    remove = set()
    for node in ast.walk(tree):
        if isinstance(node, (ast.Import, ast.ImportFrom)):
            for i in range(node.lineno - 1, node.end_lineno):
                remove.add(i)
    return "".join(line for i, line in enumerate(lines) if i not in remove)


def patch_source(src: str) -> str:
    src = src.replace('os.path.dirname(os.path.dirname(__file__))', "APP_DIR")
    src = src.replace('os.path.dirname(os.path.dirname(os.path.abspath(__file__)))', "APP_DIR")
    src = src.replace('os.path.join(base, "assets", name)', 'os.path.join(base, name)')
    src = src.replace(
        'os.path.join(os.path.dirname(os.path.dirname(__file__)), "data", "water_demand.db")',
        "DB_PATH",
    )
    src = src.replace(
        'os.path.join(os.path.dirname(os.path.dirname(__file__)), "assets", "logo.png")',
        "LOGO_PATH",
    )
    src = src.replace(
        'os.path.normpath(\n        os.path.join(os.path.dirname(os.path.dirname(__file__)), "..", "templates", "WaterDemand_Template.xlsx")\n    )',
        'os.path.join(APP_DIR, "templates", "WaterDemand_Template.xlsx")',
    )
    return src


FILES = [
    "config/nbc_2026.py",
    "config/environmental.py",
    "config/page_visibility.py",
    "models/project.py",
    "models/residential.py",
    "models/commercial.py",
    "models/other_details.py",
    "models/calculations.py",
    "models/environmental.py",
    "services/calculator.py",
    "services/database.py",
    "services/lookup_db.py",
    "services/project_service.py",
    "services/automation.py",
    "services/environmental_calculator.py",
    "services/result_tables.py",
    "services/pdf_exporter.py",
    "services/excel_exporter.py",
    "ui/components/validation.py",
    "ui/gui_safe.py",
    "ui/app_state.py",
    "ui/components/scrollable_frame.py",
    "ui/components/result_table.py",
    "ui/components/preview_dialog.py",
    "ui/splash_screen.py",
    "ui/dashboard.py",
    "ui/project_hub.py",
    "ui/project_edit_dialog.py",
    "ui/pages/project_page.py",
    "ui/pages/residential_page.py",
    "ui/pages/commercial_page.py",
    "ui/pages/other_page.py",
    "ui/pages/final_page.py",
    "_app_sidebar.py",
    "app_launcher.py",
]

parts = [HEADER]
for rel in FILES:
    path = os.path.join(BASE, rel)
    with open(path) as f:
        src = patch_source(remove_imports(f.read()))
    parts.append(f"\n# {'='*20} {rel} {'='*20}\n")
    parts.append(src)

content = "".join(parts)

# Single entry point: Application from app_launcher
if "def main():" in content:
    content = content.replace(
        'def main():\n    from app_launcher import main as launch_main\n    launch_main()',
        'def main():\n    Application().mainloop()',
        1,
    )

with open(OUT, "w") as f:
    f.write(content)

print(f"Built {OUT}: {len(content.splitlines())} lines")
